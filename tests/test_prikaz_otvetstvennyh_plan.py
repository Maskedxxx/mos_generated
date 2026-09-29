# -*- coding: utf-8 -*-
"""
Приложение № 3 «План ключевых задач» приказа 0.2 (замечание тестировщика 14.07, решение владельца 29.09):
31 задача по образцу заказчика (Джифлорс), сроки — от даты акта 0.1 по «Калькулятору дат», ответственные — по ролям.
Эталон сроков и ролей — Приложение № 3 Джифлорс (старт 29.04.2026); отличие одно — ИЦ предприятия по калькулятору (+119).
"""
import io

import pytest
from docx import Document

import engine
from generators.prikaz_otvetstvennyh import PLAN, parse_start_date, short_fio

# Тестовые ответственные (вымышленные ФИО), по ролям образца: директор, руководитель Программы, три направления
VALUES = {
    "org_full": "ООО «Тест-Флорс»", "prikaz_num": "02/05", "prikaz_date": "«14» мая 2026 г.",
    "resp1": "руководитель обособленного подразделения Петров П.П., +7 900 000-00-01, p@test.ru",
    "resp2": "руководитель экономического отдела Сидорова А.Б., +7 900 000-00-02, s@test.ru",
    "resp3": "руководитель производства Кузнецов Олег Игоревич, +7 900 000-00-03, k@test.ru",
    "resp4": "руководитель проекта по управлению изменениями Смирнова Е.Д., +7 900 000-00-04, sm@test.ru",
    "resp5": "менеджер по управлению изменениями Орлов Н.Ю., +7 900 000-00-05, o@test.ru",
    "pilot_potok": "производство паркета", "po_date": "13.07.2026",
    "resp_upt_dat": "Руководителю экономического отдела Сидоровой А.Б.",
    "resp_opt_dat": "Руководителю производства Кузнецову О.И.",
    "resp_upi_dat": "Руководителю проекта Смирновой Е.Д.", "oznak_dat": "Менеджеру Орлову Н.Ю.",
    "signer_post": "Генеральный директор", "signer_fio": "Васильев Е.В.", "act_date": "«29» апреля 2026 г.",
}
D, P, U, O, I = "Васильев Е.В.", "Петров П.П.", "Сидорова А.Б.", "Кузнецов О.И.", "Смирнова Е.Д."
# Сроки и роли Приложения № 3 Джифлорс при старте 29.04.2026 (ИЦ предприятия — по калькулятору: 26.08.2026)
EXPECTED = [
    ("29.04.2026", D), ("06.05.2026", U), ("13.05.2026", D), ("20.05.2026", U), ("20.05.2026", I),
    ("27.05.2026", D + "\n" + I), ("27.05.2026", I), ("03.06.2026", O), ("03.06.2026", O), ("24.06.2026", O),
    ("01.07.2026", I), ("08.07.2026", O), ("08.07.2026", P), ("08.07.2026", I), ("22.07.2026", O),
    ("15.07.2026", O), ("22.07.2026", O), ("22.07.2026", O), ("29.07.2026", U), ("12.08.2026", I),
    ("26.08.2026", U), ("02.09.2026", I), ("23.09.2026", O), ("14.10.2026", O), ("14.10.2026", I),
    ("21.10.2026", I), ("28.10.2026", U), ("28.10.2026", D), ("28.10.2026", O), ("28.10.2026", I),
    ("04.11.2026", D),
]


def _plan_table(data: bytes):
    """Таблица Приложения № 3 сгенерированного документа (шапка «№ | Задача | Срок (до) | Ответственный»)."""
    doc = Document(io.BytesIO(data))
    return next(t for t in doc.tables if t.rows[0].cells[-1].text.strip() == "Ответственный")


def test_plan_table_rows_dates_and_responsibles():
    """31 строка: задачи дословно из PLAN, сроки и ответственные совпадают с образцом заказчика."""
    table = _plan_table(engine.generate("prikaz_otvetstvennyh", "t", VALUES))
    rows = [[c.text for c in r.cells] for r in table.rows[1:]]
    assert len(rows) == len(PLAN) == len(EXPECTED) == 31
    for (num, task, date, resp), (p_task, _, _), (e_date, e_resp) in zip(rows, PLAN, EXPECTED):
        assert task == p_task
        assert (date, resp) == (e_date, e_resp), task


def test_responsible_runs_not_struck_through():
    """В ячейке «Ответственный» абзац шаблона зачёркнут (w:strike) — ФИО должно выводиться без зачёркивания."""
    table = _plan_table(engine.generate("prikaz_otvetstvennyh", "t", VALUES))
    for row in table.rows[1:]:
        for p in row.cells[3].paragraphs:
            for r in p.runs:
                assert not r.font.strike, row.cells[1].text


def test_no_template_tags_left():
    data = engine.generate("prikaz_otvetstvennyh", "t", VALUES)
    import zipfile
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode("utf-8")
    assert "{{" not in xml and "{%" not in xml


@pytest.mark.parametrize("text", ["«29» апреля 2026 г.", "29 апреля 2026", "29.04.2026", '"29.04.2026"', "«29» Апреля 2026г."])
def test_start_date_formats(text):
    from datetime import date
    assert parse_start_date(text) == date(2026, 4, 29)


def test_bad_start_date_is_explained():
    with pytest.raises(ValueError, match="Дата начала мероприятий не распознана"):
        engine.generate("prikaz_otvetstvennyh", "t", {**VALUES, "act_date": "апрель 2026"})


@pytest.mark.parametrize("text,expected", [
    ("руководитель обособленного подразделения Лачугин М.А., +7(929)972-15-63", "Лачугин М.А."),
    ("начальник производства, Иванов Иван Иванович, 89000000000", "Иванов И.И."),
    ("И.И. Иванов", "Иванов И.И."),
    ("Иванов И. И.", "Иванов И.И."),
    ("директор по персоналу Кузнецова-Лебедь Анна Сергеевна", "Кузнецова-Лебедь А.С."),
    ("менеджер, телефон 8900", None),
])
def test_short_fio(text, expected):
    assert short_fio(text) == expected


def test_responsible_without_fio_is_explained():
    """Поле роли без ФИО — ошибка с подписью поля, а не пустая/ложная ячейка в плане."""
    with pytest.raises(ValueError, match="Ответственный 3 — «Оптимизация потоков».*не найдено ФИО"):
        engine.generate("prikaz_otvetstvennyh", "t", {**VALUES, "resp3": "начальник цеха, тел. 8900"})


def test_missing_required_lists_labels_before_plan_errors():
    """Пустая форма — привычный список незаполненных полей (включая дату начала), а не ошибка разбора ФИО."""
    with pytest.raises(ValueError, match="Не заполнены обязательные поля:.*Дата начала мероприятий"):
        engine.generate("prikaz_otvetstvennyh", "t", {})


def test_start_date_prefilled_from_act_via_org_store():
    """Дата акта 0.1 сохраняется в стор по ООО под ключом act_date — форма 0.2 получает её подсказкой по тому же ООО."""
    from conftest import sample_values
    from registry import GENERATORS
    act = {**sample_values(GENERATORS["akt_nachala"]().SCHEMA), "org_full": "ООО «Тест-Флорс»", "act_date": "«29» апреля 2026 г."}
    engine.generate("akt_nachala", "t", act)
    assert engine.org_suggest("Тест-Флорс")["act_date"][0] == "«29» апреля 2026 г."
    assert any(f["key"] == "act_date" for f in GENERATORS["prikaz_otvetstvennyh"]().SCHEMA)
